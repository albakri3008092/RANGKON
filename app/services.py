import csv
import io
import re
from collections.abc import Iterable
from datetime import datetime, timedelta

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session, joinedload, sessionmaker

from .config import settings
from .models import STATUS_RANK, Event, Invitation, Member
from .whatsapp import Provider, SendError

DEFAULT_MESSAGE = (
    "Salam {nama} ({unit}),\n\n"
    "Anda dipanggil untuk berkumpul:\n"
    "*{tajuk}*\n"
    "Tempat: {lokasi}\n"
    "Masa: {masa}\n\n"
    "WAJIB balas *TERIMA* untuk mengesahkan panggilan ini diterima "
    "(atau *TOLAK* jika tidak dapat hadir).\n"
    "Atau tekan pautan: {pautan}"
)

ACCEPT_WORDS = {
    "terima",
    "diterima",
    "accept",
    "ya",
    "yes",
    "y",
    "ok",
    "okay",
    "hadir",
    "baik",
    "roger",
    "sedia",
}
DECLINE_WORDS = {"tolak", "decline", "tidak", "tak", "no", "n", "tdk", "takleh", "uzur", "cuti"}
DECLINE_PHRASES = ("tak boleh", "tidak boleh", "tidak hadir", "tak hadir", "tidak dapat")


def now() -> datetime:
    return datetime.now(settings.timezone).replace(tzinfo=None)


def from_unix(ts: str | int | None) -> datetime:
    if ts is None:
        return now()
    return datetime.fromtimestamp(int(ts), settings.timezone).replace(tzinfo=None)


def normalize_phone(raw: str) -> str | None:
    digits = re.sub(r"\D", "", raw or "")
    if not digits:
        return None
    if digits.startswith("00"):
        digits = digits[2:]
    elif digits.startswith("0"):
        digits = "60" + digits[1:]
    elif digits.startswith("1") and 9 <= len(digits) <= 10:
        digits = "60" + digits
    if not 10 <= len(digits) <= 15:
        return None
    return digits


def parse_rsvp(text: str) -> str | None:
    cleaned = " ".join(re.sub(r"[^\w\s]", " ", (text or "").lower()).split())
    if not cleaned:
        return None
    if any(p in cleaned for p in DECLINE_PHRASES):
        return "declined"
    first = cleaned.split()[0]
    if first in DECLINE_WORDS:
        return "declined"
    if first in ACCEPT_WORDS:
        return "accepted"
    return None


def set_rsvp(inv: Invitation, answer: str, via: str, at: datetime | None = None) -> None:
    inv.rsvp = answer
    inv.rsvp_at = at or now()
    inv.rsvp_via = via


RSVP_LABEL = {"accepted": "Terima", "declined": "Tolak"}


def format_dt(dt: datetime | None) -> str:
    if dt is None:
        return ""
    return dt.strftime("%d/%m/%Y %I:%M %p")


def checkin_url(inv: Invitation) -> str:
    return f"{settings.public_base_url}/c/{inv.token}"


def message_params(inv: Invitation) -> dict[str, str]:
    m, e = inv.member, inv.event
    nama = f"{m.rank} {m.name}".strip()
    return {
        "nama": nama,
        "tajuk": e.title,
        "lokasi": e.location,
        "masa": format_dt(e.starts_at),
        "pautan": checkin_url(inv),
        "unit": inv.unit,
    }


def render_message(template: str, params: dict[str, str]) -> str:
    out = template
    for key, value in params.items():
        out = out.replace("{" + key + "}", value)
    return out


def checkin_window_open(event: Event, at: datetime | None = None) -> bool:
    at = at or now()
    return at >= event.starts_at - timedelta(minutes=settings.checkin_opens_minutes_before)


def mark_attended(inv: Invitation, method: str) -> bool:
    if inv.attended_at is not None:
        return False
    inv.attended_at = now()
    inv.checkin_method = method
    return True


# --- members -----------------------------------------------------------------

CSV_ALIASES = {
    "name": {"nama", "name", "nama penuh"},
    "phone": {"telefon", "phone", "no telefon", "no. telefon", "nombor telefon", "hp", "no hp"},
    "unit": {"unit", "kumpulan", "group", "pasukan"},
    "rank": {"pangkat", "rank"},
    "service_no": {"no tentera", "no. tentera", "no perkhidmatan", "service no", "no siri"},
}


def match_unit(raw: str) -> str | None:
    key = re.sub(r"[\s.]", "", (raw or "").lower())
    for unit in settings.units:
        if re.sub(r"[\s.]", "", unit.lower()) == key:
            return unit
    return None


def xlsx_to_csv(raw: bytes) -> str:
    sheet = load_workbook(io.BytesIO(raw), read_only=True, data_only=True).active
    buf = io.StringIO()
    writer = csv.writer(buf)
    for row in sheet.iter_rows(values_only=True):
        writer.writerow(["" if v is None else str(v) for v in row])
    return buf.getvalue()


def import_members_csv(db: Session, content: str, default_unit: str | None) -> dict[str, int]:
    reader = csv.DictReader(io.StringIO(content))
    colmap: dict[str, str] = {}
    for col in reader.fieldnames or []:
        norm = col.strip().lower()
        for field, aliases in CSV_ALIASES.items():
            if norm in aliases:
                colmap[field] = col
    result = {"added": 0, "updated": 0, "skipped": 0}
    if "name" not in colmap or "phone" not in colmap:
        result["skipped"] = -1
        return result
    for row in reader:
        name = (row.get(colmap["name"]) or "").strip()
        phone = normalize_phone(row.get(colmap["phone"]) or "")
        unit = match_unit(row.get(colmap["unit"], "") if "unit" in colmap else "") or match_unit(
            default_unit or ""
        )
        if not name or not phone or not unit:
            result["skipped"] += 1
            continue
        rank = (row.get(colmap["rank"]) or "").strip() if "rank" in colmap else ""
        service_no = (row.get(colmap["service_no"]) or "").strip() if "service_no" in colmap else ""
        member = db.scalar(select(Member).where(Member.phone == phone))
        if member:
            member.name, member.unit = name, unit
            if rank:
                member.rank = rank
            if service_no:
                member.service_no = service_no
            result["updated"] += 1
        else:
            db.add(
                Member(
                    name=name,
                    phone=phone,
                    unit=unit,
                    rank=rank,
                    service_no=service_no,
                    created_at=now(),
                )
            )
            result["added"] += 1
        db.flush()
    db.commit()
    return result


# --- invitations -------------------------------------------------------------


def create_invitations(db: Session, event: Event, units: Iterable[str]) -> list[int]:
    units = list(units)
    existing = set(db.scalars(select(Invitation.member_id).where(Invitation.event_id == event.id)))
    members = db.scalars(
        select(Member).where(Member.active.is_(True), Member.unit.in_(units))
    ).all()
    new = [
        Invitation(event_id=event.id, member_id=m.id, unit=m.unit, phone=m.phone)
        for m in members
        if m.id not in existing
    ]
    db.add_all(new)
    db.commit()
    return [inv.id for inv in new]


def deliver(factory: sessionmaker, invitation_ids: list[int], provider: Provider) -> None:
    with factory() as db:
        for inv_id in invitation_ids:
            inv = db.scalar(
                select(Invitation)
                .options(joinedload(Invitation.member), joinedload(Invitation.event))
                .where(Invitation.id == inv_id)
            )
            if inv is None or inv.status not in ("pending", "failed"):
                continue
            params = message_params(inv)
            body = render_message(inv.event.message, params)
            try:
                inv.phone = inv.member.phone
                inv.wa_message_id = provider.send(inv.phone, body, params)
                inv.status = "sent"
                inv.error = None
                inv.sent_at = now()
            except SendError as exc:
                inv.status = "failed"
                inv.error = str(exc)[:500]
            db.commit()


def apply_status(inv: Invitation, status: str, at: datetime, error: str | None = None) -> None:
    if status == "failed":
        if inv.status in ("pending", "sent"):
            inv.status = "failed"
            inv.error = error or "Gagal dihantar"
        return
    if status not in STATUS_RANK:
        return
    if STATUS_RANK[status] > STATUS_RANK.get(inv.status, 0) or inv.status == "failed":
        inv.status = status
        inv.error = None
    if status in ("delivered", "read") and inv.delivered_at is None:
        inv.delivered_at = at
    if status == "read" and inv.read_at is None:
        inv.read_at = at


def record_reply(inv: Invitation, text: str, at: datetime) -> None:
    apply_status(inv, "read", at)
    if inv.replied_at and at < inv.replied_at:
        return
    inv.reply_text = text[:1000]
    inv.replied_at = at
    answer = parse_rsvp(text)
    if answer and not (inv.rsvp_at and at < inv.rsvp_at):
        set_rsvp(inv, answer, "whatsapp", at)


def handle_webhook(db: Session, payload: dict) -> dict[str, int]:
    counts = {"statuses": 0, "replies": 0}
    for entry in payload.get("entry", []):
        for change in entry.get("changes", []):
            value = change.get("value", {})
            for st in value.get("statuses", []):
                inv = db.scalar(select(Invitation).where(Invitation.wa_message_id == st.get("id")))
                if inv is None:
                    continue
                errors = st.get("errors") or []
                err = errors[0].get("title") if errors else None
                apply_status(inv, st.get("status", ""), from_unix(st.get("timestamp")), err)
                counts["statuses"] += 1
            for msg in value.get("messages", []):
                inv = _invitation_for_reply(db, msg)
                if inv is None:
                    continue
                text = (
                    msg.get("text", {}).get("body")
                    or msg.get("button", {}).get("text")
                    or msg.get("interactive", {}).get("button_reply", {}).get("title")
                    or ""
                )
                record_reply(inv, text, from_unix(msg.get("timestamp")))
                counts["replies"] += 1
    db.commit()
    return counts


def _invitation_for_reply(db: Session, msg: dict) -> Invitation | None:
    context_id = (msg.get("context") or {}).get("id")
    if context_id:
        inv = db.scalar(select(Invitation).where(Invitation.wa_message_id == context_id))
        if inv:
            return inv
    phone = normalize_phone(msg.get("from", ""))
    if not phone:
        return None
    return db.scalar(
        select(Invitation)
        .where(Invitation.phone == phone, Invitation.sent_at.is_not(None))
        .order_by(Invitation.sent_at.desc())
        .limit(1)
    )


# --- dashboard ---------------------------------------------------------------


def _empty_stats() -> dict[str, int]:
    return {
        "total": 0,
        "pending": 0,
        "failed": 0,
        "sent": 0,
        "delivered": 0,
        "read": 0,
        "accepted": 0,
        "declined": 0,
        "no_response": 0,
        "attended": 0,
    }


def _add(stats: dict[str, int], inv: Invitation) -> None:
    stats["total"] += 1
    if inv.status == "pending":
        stats["pending"] += 1
    elif inv.status == "failed":
        stats["failed"] += 1
    else:
        stats["sent"] += 1
        if STATUS_RANK[inv.status] >= STATUS_RANK["delivered"]:
            stats["delivered"] += 1
        if inv.status == "read":
            stats["read"] += 1
    if inv.rsvp == "accepted":
        stats["accepted"] += 1
    elif inv.rsvp == "declined":
        stats["declined"] += 1
    else:
        stats["no_response"] += 1
    if inv.attended_at:
        stats["attended"] += 1


def unit_sort_key(unit: str) -> tuple[int, str]:
    try:
        return (settings.units.index(unit), unit)
    except ValueError:
        return (len(settings.units), unit)


def event_dashboard(db: Session, event: Event) -> dict:
    invitations = db.scalars(
        select(Invitation)
        .options(joinedload(Invitation.member))
        .where(Invitation.event_id == event.id)
    ).all()
    overall = _empty_stats()
    by_unit = {u: _empty_stats() for u in settings.units}
    rows = []
    for inv in sorted(invitations, key=lambda i: (unit_sort_key(i.unit), i.member.name.lower())):
        _add(overall, inv)
        _add(by_unit.setdefault(inv.unit, _empty_stats()), inv)
        rows.append(
            {
                "id": inv.id,
                "name": inv.member.name,
                "rank": inv.member.rank,
                "phone": inv.phone,
                "unit": inv.unit,
                "status": inv.status,
                "error": inv.error,
                "rsvp": inv.rsvp,
                "rsvp_at": format_dt(inv.rsvp_at),
                "rsvp_via": inv.rsvp_via,
                "reply": inv.reply_text,
                "sent_at": format_dt(inv.sent_at),
                "delivered_at": format_dt(inv.delivered_at),
                "read_at": format_dt(inv.read_at),
                "attended_at": format_dt(inv.attended_at),
                "checkin_method": inv.checkin_method,
            }
        )
    return {"overall": overall, "units": by_unit, "rows": rows}


def export_csv(dashboard: dict) -> str:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(
        [
            "Unit",
            "Pangkat",
            "Nama",
            "Telefon",
            "Status WhatsApp",
            "Penerimaan",
            "Masa Terima",
            "Kaedah Terima",
            "Balasan",
            "Dihantar",
            "Sampai",
            "Dibaca",
            "Hadir",
            "Kaedah Hadir",
        ]
    )
    for r in dashboard["rows"]:
        w.writerow(
            [
                r["unit"],
                r["rank"],
                r["name"],
                r["phone"],
                r["status"],
                RSVP_LABEL.get(r["rsvp"], "Belum terima"),
                r["rsvp_at"],
                r["rsvp_via"] or "",
                r["reply"] or "",
                r["sent_at"],
                r["delivered_at"],
                r["read_at"],
                r["attended_at"],
                r["checkin_method"] or "",
            ]
        )
    return buf.getvalue()
